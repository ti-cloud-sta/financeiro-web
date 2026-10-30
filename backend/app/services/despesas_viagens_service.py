from sqlalchemy.orm import Session
from fastapi import HTTPException, UploadFile
import io
import openpyxl
from datetime import datetime
from app.models.movimentacao import Movimentacao
from app.models.importacao import Importacao
from app.repositories.categoria_repository import CategoriaRepository
from app.repositories.colaborador_repository import ColaboradorRepository
from app.repositories.empresa_repository import EmpresaRepository
from app.repositories.colaborador_alias_repository import ColaboradorAliasRepository
from app.repositories.categoria_alias_repository import CategoriaAliasRepository
from app.schemas.despesas_viagens import (
    SalvarDespesaViagemPayload,
    ConciliacaoViagensResponse,
    ConciliacaoViagemUnidadeRow,
    ConciliacaoViagemEmpresaRow,
    ConciliacaoViagemPessoaRow
)


class DespesasViagensService:
    def __init__(self, db: Session):
        self.db = db
        self.cat_repo = CategoriaRepository(db)
        self.colab_repo = ColaboradorRepository(db)
        self.emp_repo = EmpresaRepository(db)
        self.alias_repo = ColaboradorAliasRepository(db)
        self.categoria_alias_repo = CategoriaAliasRepository(db)

    def salvar_importacao(self, payload: SalvarDespesaViagemPayload):
        # Resolve a data de competência do form (criará as movimentações com essa data)
        data_competencia_obj = None
        if payload.dataCompetencia:
            try:
                data_competencia_obj = datetime.strptime(payload.dataCompetencia, "%Y-%m-%d")
            except ValueError:
                pass

        nova_importacao = None

        # -----------------------------------------------------------------------
        # Lançamento: cria um registro de importação (seja IA ou Manual)
        # -----------------------------------------------------------------------
        empresa_manual = None
        if payload.isManualEntry:
            if payload.idEmpresaManual:
                empresa_manual = self.emp_repo.get_by_id(payload.idEmpresaManual)
            elif payload.despesas and payload.despesas[0].empresa:
                empresa_manual = self.emp_repo.get_by_nome(payload.despesas[0].empresa)
            if not empresa_manual:
                raise HTTPException(status_code=400, detail="Empresa / Origem não selecionada ou não encontrada para a despesa manual.")

        extensao = payload.nomeArquivo.split('.')[-1] if '.' in payload.nomeArquivo else 'MANUAL'
        tipo_importacao = "MANUAL_DESPESAS" if payload.isManualEntry else "IA_DESPESAS"

        id_empresa_importacao = None
        if payload.isManualEntry:
            id_empresa_importacao = empresa_manual.idEmpresas
        elif payload.despesas:
            emp = self.emp_repo.get_by_nome(payload.despesas[0].empresa)
            if emp:
                id_empresa_importacao = emp.idEmpresas

        nova_importacao = Importacao(
            nomeArquivo=payload.nomeArquivo,
            extensaoArquivo=extensao,
            idEmpresa=id_empresa_importacao,
            tipo=tipo_importacao,
            idUserInc=payload.idUserInc,
            createdAt=data_competencia_obj or datetime.now()
        )
        self.db.add(nova_importacao)
        self.db.flush()

        # -----------------------------------------------------------------------
        # Grava as movimentações
        # -----------------------------------------------------------------------
        for d in payload.despesas:
            cat = self.cat_repo.get_by_nome(d.categoria)
            if not cat:
                raise HTTPException(status_code=400, detail=f"Categoria não encontrada: {d.categoria}")

            # Se o nome da categoria extraído do documento diverge do nome cadastrado,
            # registra/atualiza o alias para que próximas importações sejam resolvidas automaticamente
            categoria_original = (d.categoria_original or d.categoria).strip()
            if categoria_original.lower() != cat.nome.lower():
                self.categoria_alias_repo.create_or_update(cat.idCategorias, categoria_original)

            if payload.isManualEntry:
                emp = empresa_manual
            else:
                emp = self.emp_repo.get_by_nome(d.empresa)
            if not emp:
                raise HTTPException(status_code=400, detail=f"Empresa não encontrada: {d.empresa}")

            # nome_original = nome divergente como estava no documento (ex: "JOSE DA SILVA")
            # nome_corrigido = nome que o usuário escolheu no select (ex: "José da Silva")
            nome_original = (d.colaborador_original or d.colaborador).strip()
            nome_corrigido = d.colaborador.strip()
            colab = None

            # 1) O nome corrigido (o que o usuário confirmou/selecionou na interface) é soberano.
            # Tenta encontrar primeiro por esse nome.
            colab = self.colab_repo.get_by_nome_normalizado(nome_corrigido)

            # 2) Se não achou pelo nome corrigido, e o usuário não alterou o nome (corrigido == original),
            # verifica se já existe um alias mapeado para esse nome.
            if not colab and nome_corrigido == nome_original:
                alias = self.alias_repo.get_by_nome_divergente(nome_original)
                if alias:
                    colab = self.colab_repo.get_by_id(alias.idColaborador)

            if not colab:
                raise HTTPException(
                    status_code=400,
                    detail=f"Colaborador não encontrado e não mapeado: {d.colaborador}"
                )

            # 3) Se o nome original diverge do nome cadastrado, registra/atualiza o alias
            #    para que próximas importações com esse mesmo nome sejam resolvidas automaticamente
            if nome_original.lower() != colab.nome.lower():
                self.alias_repo.create_or_update(colab.idColaborador, nome_original)

            # A data de competência definida no modal é SEMPRE o critério de agrupamento.
            # A data interna do documento (d.data) é apenas informativa e não deve
            # sobrescrever a competência contábil definida pelo usuário na importação.
            data_mov = data_competencia_obj

            nova_mov = Movimentacao(
                idCategoria=cat.idCategorias,
                idColaborador=colab.idColaborador,
                idEmpresa=emp.idEmpresas,
                idUnidade=payload.idUnidade,
                idImportacoes=nova_importacao.idImportacoes if nova_importacao else None,
                nroDocumento=d.nroDocumento or None,
                valor=d.valor,
                createdAt=data_mov or datetime.now()
            )
            self.db.add(nova_mov)

        self.db.commit()

        return {
            "sucesso": True,
            "idImportacoes": nova_importacao.idImportacoes if nova_importacao else None,
            "totalMovimentacoes": len(payload.despesas)
        }

    def obter_visao_geral(self, filtros: dict):
        from sqlalchemy import func, extract, case
        from datetime import datetime, timedelta
        from app.models.movimentacao import Movimentacao
        from app.models.importacao import Importacao
        from app.models.colaborador import Colaborador
        from app.models.empresa import Empresa
        from app.models.categoria import Categoria
        from app.models.unidade import Unidade

        db = self.db

        # ---- Parse filtros ----
        data_inicio = filtros.get("data_inicio")
        data_fim = filtros.get("data_fim")
        id_empresa = filtros.get("id_empresa")
        id_colaborador = filtros.get("id_colaborador")
        id_categoria = filtros.get("id_categoria")
        id_unidade = filtros.get("id_unidade")

        def parse_date(d):
            if not d:
                return None
            for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
                try:
                    return datetime.strptime(d, fmt)
                except ValueError:
                    pass
            return None

        dt_inicio = parse_date(data_inicio)
        dt_fim = parse_date(data_fim)

        # ---- Base query ----
        def base_q():
            q = db.query(Movimentacao)
            if dt_inicio:
                q = q.filter(Movimentacao.createdAt >= dt_inicio)
            if dt_fim:
                # inclui o dia todo
                q = q.filter(Movimentacao.createdAt <= dt_fim.replace(hour=23, minute=59, second=59))
            if id_empresa:
                q = q.filter(Movimentacao.idEmpresa == id_empresa)
            if id_colaborador:
                q = q.filter(Movimentacao.idColaborador == id_colaborador)
            if id_categoria:
                q = q.filter(Movimentacao.idCategoria == id_categoria)
            if id_unidade:
                q = q.filter(Movimentacao.idUnidade == id_unidade)
            # Apenas importações do tipo IA_DESPESAS ou MANUAL_DESPESAS
            q = q.join(Importacao, Movimentacao.idImportacoes == Importacao.idImportacoes)
            q = q.filter(Importacao.tipo.in_(["IA_DESPESAS", "MANUAL_DESPESAS", "Despesa Manual"]))
            return q

        # ---- KPI: Total e Quantidade ----
        agg = base_q().with_entities(
            func.coalesce(func.sum(Movimentacao.valor), 0).label("total"),
            func.count(Movimentacao.idMovimentacoes).label("qtd")
        ).one()
        total = float(agg.total)
        qtd = int(agg.qtd)

        # ---- KPI: Ticket Médio ----
        ticket_medio = round(total / qtd, 2) if qtd > 0 else 0.0

        # ---- KPI: Total do Mês atual ----
        agora = datetime.now()
        inicio_mes = agora.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        # Mes anterior (ultimo dia do mes anterior)
        inicio_mes_ant = (inicio_mes - timedelta(days=1)).replace(day=1, hour=0, minute=0, second=0)
        fim_mes_ant = inicio_mes - timedelta(seconds=1)

        def soma_mes(inicio, fim):
            q = db.query(func.coalesce(func.sum(Movimentacao.valor), 0))
            q = q.join(Importacao, Movimentacao.idImportacoes == Importacao.idImportacoes)
            q = q.filter(Importacao.tipo.in_(["IA_DESPESAS", "MANUAL_DESPESAS", "Despesa Manual"]))
            q = q.filter(Movimentacao.createdAt >= inicio, Movimentacao.createdAt <= fim)
            if id_empresa:
                q = q.filter(Movimentacao.idEmpresa == id_empresa)
            if id_colaborador:
                q = q.filter(Movimentacao.idColaborador == id_colaborador)
            if id_categoria:
                q = q.filter(Movimentacao.idCategoria == id_categoria)
            if id_unidade:
                q = q.filter(Movimentacao.idUnidade == id_unidade)
            return float(q.scalar() or 0)

        total_mes = soma_mes(inicio_mes, agora)
        total_mes_ant = soma_mes(inicio_mes_ant, fim_mes_ant)
        if total_mes_ant > 0:
            pct_mes = round(((total_mes - total_mes_ant) / total_mes_ant) * 100, 1)
        elif total_mes > 0:
            pct_mes = 100.0
        else:
            pct_mes = 0.0

        # ---- KPI: Maior Despesa ----
        maior = base_q().with_entities(
            Movimentacao.valor,
            Movimentacao.idColaborador,
            Movimentacao.idCategoria,
            Movimentacao.idUnidade
        ).order_by(Movimentacao.valor.desc()).first()

        maior_valor = 0.0
        maior_contexto = "Sem registros"
        if maior:
            maior_valor = float(maior.valor)
            colab = db.query(Colaborador).filter(Colaborador.idColaborador == maior.idColaborador).first()
            cat = db.query(Categoria).filter(Categoria.idCategorias == maior.idCategoria).first()
            uni = db.query(Unidade).filter(Unidade.idUnidade == maior.idUnidade).first() if maior.idUnidade else None
            nome_colab = colab.nome if colab else "?"
            nome_cat = cat.nome if cat else "?"
            nome_uni = uni.descricao if uni else "N/D"
            maior_contexto = f"{nome_colab} · {nome_uni} · {nome_cat}"

        # ---- Tabela Maiores Despesas (top 10) ----
        top_rows = (
            base_q()
            .join(Colaborador, Movimentacao.idColaborador == Colaborador.idColaborador)
            .join(Empresa, Movimentacao.idEmpresa == Empresa.idEmpresas)
            .join(Categoria, Movimentacao.idCategoria == Categoria.idCategorias)
            .with_entities(
                Movimentacao.nroDocumento,
                Colaborador.nome.label("colaborador"),
                Empresa.nome.label("empresa"),
                Categoria.nome.label("categoria"),
                Movimentacao.createdAt,
                Movimentacao.valor
            )
            .order_by(Movimentacao.valor.desc())
            .limit(10)
            .all()
        )
        tabela_maiores = [
            {
                "descricao": r.nroDocumento or "—",
                "pessoa": r.colaborador,
                "empresa": r.empresa,
                "categoria": r.categoria,
                "data": r.createdAt.strftime("%d/%m/%Y") if r.createdAt else "—",
                "valor": float(r.valor)
            }
            for r in top_rows
        ]

        # ---- Donut Categorias ----
        cat_agg = (
            base_q()
            .join(Categoria, Movimentacao.idCategoria == Categoria.idCategorias)
            .with_entities(
                Categoria.idCategorias.label("id"),
                Categoria.nome.label("nome"),
                func.sum(Movimentacao.valor).label("total"),
                func.count(Movimentacao.idMovimentacoes).label("qtd")
            )
            .group_by(Categoria.idCategorias, Categoria.nome)
            .order_by(func.sum(Movimentacao.valor).desc())
            .limit(8)
            .all()
        )
        donut_categorias = [{"id": r.id, "name": r.nome, "value": round(float(r.total), 2), "qtd": r.qtd} for r in cat_agg]

        # ---- Donut Empresas ----
        emp_agg = (
            base_q()
            .join(Empresa, Movimentacao.idEmpresa == Empresa.idEmpresas)
            .with_entities(
                Empresa.nome.label("nome"),
                func.sum(Movimentacao.valor).label("total")
            )
            .group_by(Empresa.nome)
            .order_by(func.sum(Movimentacao.valor).desc())
            .all()
        )
        donut_empresas = [{"name": r.nome, "value": round(float(r.total), 2)} for r in emp_agg]

        # ---- Evolução Mensal por Categoria ----
        # 1. Determina a janela de meses (mesmo os sem dados)
        min_max = base_q().with_entities(func.min(Movimentacao.createdAt), func.max(Movimentacao.createdAt)).one()
        min_dt = dt_inicio or min_max[0]
        max_dt = dt_fim or min_max[1]
        
        all_months = []
        if min_dt and max_dt:
            curr = min_dt.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
            end = max_dt.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
            while curr <= end:
                all_months.append((curr.year, curr.month))
                month = curr.month % 12 + 1
                year = curr.year + (1 if curr.month == 12 else 0)
                curr = curr.replace(year=year, month=month)
                
        meses_labels = [f"{m:02d}/{a}" for a, m in all_months]

        # 2. Top 5 categorias por total
        top_cats = [r.nome for r in cat_agg[:5]]

        # 3. Pega os dados brutos de evolucao da query para as top categorias (para evitar N queries no loop)
        evolucao_raw = (
            base_q()
            .join(Categoria, Movimentacao.idCategoria == Categoria.idCategorias)
            .filter(Categoria.nome.in_(top_cats) if top_cats else False)
            .with_entities(
                Categoria.nome.label("cat_nome"),
                func.year(Movimentacao.createdAt).label("ano"),
                func.month(Movimentacao.createdAt).label("mes"),
                func.sum(Movimentacao.valor).label("total")
            )
            .group_by(Categoria.nome, func.year(Movimentacao.createdAt), func.month(Movimentacao.createdAt))
            .all()
        )
        
        # Mapa: dict[cat_nome][ano][mes] = total
        evolucao_dict = {}
        for r in evolucao_raw:
            if r.cat_nome not in evolucao_dict:
                evolucao_dict[r.cat_nome] = {}
            if r.ano not in evolucao_dict[r.cat_nome]:
                evolucao_dict[r.cat_nome][r.ano] = {}
            evolucao_dict[r.cat_nome][r.ano][r.mes] = float(r.total)

        series = []
        for cat_nome in top_cats:
            data_points = []
            for ano, mes in all_months:
                val = evolucao_dict.get(cat_nome, {}).get(ano, {}).get(mes, 0.0)
                data_points.append(round(val, 2))
            series.append({"name": cat_nome, "data": data_points})

        # ---- Mapa: Despesas por Estado ----
        from app.models.centro_custo import CentroEstado
        
        mapa_agg = (
            base_q()
            .join(Colaborador, Movimentacao.idColaborador == Colaborador.idColaborador)
            .join(CentroEstado, Colaborador.idCentroCusto == CentroEstado.idCentroCusto)
            .with_entities(
                CentroEstado.estado.label("estado"),
                func.sum(Movimentacao.valor).label("total"),
                func.count(Movimentacao.idMovimentacoes).label("qtd")
            )
            .group_by(CentroEstado.estado)
            .all()
        )
        
        # Mapeamento do nome completo do estado caso esteja em sigla (Ex: 'SP' -> 'São Paulo')
        # O echarts brazil map precisa do nome exato e case-sensitive do estado.
        sigla_para_nome = {
            "AC": "Acre", "AL": "Alagoas", "AP": "Amapá", "AM": "Amazonas", "BA": "Bahia",
            "CE": "Ceará", "DF": "Distrito Federal", "ES": "Espírito Santo", "GO": "Goiás",
            "MA": "Maranhão", "MT": "Mato Grosso", "MS": "Mato Grosso do Sul", "MG": "Minas Gerais",
            "PA": "Pará", "PB": "Paraíba", "PR": "Paraná", "PE": "Pernambuco", "PI": "Piauí",
            "RJ": "Rio de Janeiro", "RN": "Rio Grande do Norte", "RS": "Rio Grande do Sul",
            "RO": "Rondônia", "RR": "Roraima", "SC": "Santa Catarina", "SP": "São Paulo",
            "SE": "Sergipe", "TO": "Tocantins"
        }
        
        # Permite resolver tanto se vier "SP" quanto se vier "SÃO PAULO" do banco
        upper_to_correct = {v.upper(): v for v in sigla_para_nome.values()}
        upper_to_correct.update(sigla_para_nome)
        
        mapa_data = []
        for r in mapa_agg:
            estado_original = (r.estado or "").strip().upper()
            nome_estado = upper_to_correct.get(estado_original, estado_original)
            mapa_data.append({
                "name": nome_estado,
                "value": round(float(r.total), 2),
                "qtd": int(r.qtd)
            })
        # ---- Spenders (Maiores Gastos por Pessoa) ----
        spenders_agg = (
            base_q()
            .join(Colaborador, Movimentacao.idColaborador == Colaborador.idColaborador)
            .with_entities(
                Colaborador.nome.label("nome"),
                func.sum(Movimentacao.valor).label("total")
            )
            .group_by(Colaborador.nome)
            .order_by(func.sum(Movimentacao.valor).desc())
            .limit(10)
            .all()
        )
        spenders = [{"name": r.nome, "value": round(float(r.total), 2)} for r in spenders_agg]

        return {
            "dashVisaoGeral": {
                "total": round(total, 2),
                "quantidadeDespesas": qtd,
                "totalMes": round(total_mes, 2),
                "percentualMes": pct_mes,
                "ticketMedio": round(ticket_medio, 2),
                "ticketMedioPercentual": 0.0,
                "maiorDespesa": round(maior_valor, 2),
                "maiorDespesaContexto": maior_contexto
            },
            "tabelaMaioresDespesas": tabela_maiores,
            "donutCategorias": donut_categorias,
            "donutEmpresas": donut_empresas,
            "mapaData": mapa_data,
            "spenders": spenders,
            "evolucao": {
                "meses": meses_labels,
                "series": series
            }
        }

    def obter_visao_comercial(self, filtros):
        from sqlalchemy import func
        from datetime import datetime
        from app.models.movimentacao import Movimentacao
        from app.models.importacao import Importacao
        from app.models.colaborador import Colaborador
        from app.models.centro_custo import CentroCusto
        from app.models.categoria import Categoria

        data_inicio_str = filtros.get("data_inicio")
        data_fim_str = filtros.get("data_fim")
        id_empresa = filtros.get("id_empresa")
        id_colaborador = filtros.get("id_colaborador")
        id_categoria = filtros.get("id_categoria")
        id_unidade = filtros.get("id_unidade")

        # Base query for all Despesas de Viagens
        def base_q():
            q = self.db.query(Movimentacao)
            q = q.join(Importacao, Movimentacao.idImportacoes == Importacao.idImportacoes)
            q = q.filter(Importacao.tipo.in_(["IA_DESPESAS", "MANUAL_DESPESAS", "Despesa Manual"]))
            if data_inicio_str:
                di = datetime.strptime(data_inicio_str, "%Y-%m-%d")
                q = q.filter(Movimentacao.createdAt >= di)
            if data_fim_str:
                df = datetime.strptime(data_fim_str, "%Y-%m-%d").replace(hour=23, minute=59, second=59)
                q = q.filter(Movimentacao.createdAt <= df)
            if id_empresa:
                q = q.filter(Movimentacao.idEmpresa == id_empresa)
            if id_colaborador:
                q = q.filter(Movimentacao.idColaborador == id_colaborador)
            if id_categoria:
                q = q.filter(Movimentacao.idCategoria == id_categoria)
            if id_unidade:
                q = q.filter(Movimentacao.idUnidade == id_unidade)
            
            # Filtro Comercial/Marketing
            q = q.join(Colaborador, Movimentacao.idColaborador == Colaborador.idColaborador)
            q = q.filter(
                Colaborador.papel.ilike("%COMERCIAL%") |
                Colaborador.papel.ilike("%MARKETING%")
            )
            return q

        total = base_q().with_entities(func.coalesce(func.sum(Movimentacao.valor), 0)).scalar()
        total = float(total or 0)

        movs = base_q().with_entities(
            Movimentacao.valor,
            Movimentacao.createdAt,
            Categoria.nome.label("categoria_nome"),
            Colaborador.nome.label("colab_nome"),
            Colaborador.papel.label("colab_papel"),
            CentroCusto.nome.label("cc_nome"),
            CentroCusto.codigo.label("cc_codigo")
        ).outerjoin(Categoria, Movimentacao.idCategoria == Categoria.idCategorias) \
         .outerjoin(CentroCusto, Colaborador.idCentroCusto == CentroCusto.codigo).all()

        meses_set = set()
        barras_vert_dict = {}
        cc_evolucao = {}
        cat_barras_dict = {}
        butterfly_cats = set()
        comercial_dict = {}
        marketing_dict = {}
        cc_barras_dict = {}
        ranking_colab_dict = {}
        ranking_cat_dict = {}
        matriz_colabs = {}

        for r in movs:
            val = float(r.valor or 0)
            dt = r.createdAt
            if not dt: continue
            mes_key = f"{dt.year}-{dt.month:02d}"
            
            cat = r.categoria_nome or "Sem Categoria"
            colab = r.colab_nome or "Desconhecido"
            cc = f"{r.cc_codigo} - {r.cc_nome}"
            colab_papel = (r.colab_papel or "").upper()

            meses_set.add(mes_key)
            barras_vert_dict[mes_key] = barras_vert_dict.get(mes_key, 0) + val
            
            if cc not in cc_evolucao: cc_evolucao[cc] = {}
            cc_evolucao[cc][mes_key] = cc_evolucao[cc].get(mes_key, 0) + val
            
            cat_barras_dict[cat] = cat_barras_dict.get(cat, 0) + val
            cc_barras_dict[cc] = cc_barras_dict.get(cc, 0) + val
            
            ranking_colab_dict[colab] = ranking_colab_dict.get(colab, 0) + val
            ranking_cat_dict[cat] = ranking_cat_dict.get(cat, 0) + val
            
            butterfly_cats.add(cat)
            if "MARKETING" in colab_papel:
                marketing_dict[cat] = marketing_dict.get(cat, 0) + val
            else:
                comercial_dict[cat] = comercial_dict.get(cat, 0) + val

            if colab not in matriz_colabs:
                matriz_colabs[colab] = {"total": 0, "valores": {}}
            matriz_colabs[colab]["total"] += val
            matriz_colabs[colab]["valores"][cat] = matriz_colabs[colab]["valores"].get(cat, 0) + val

        sorted_mes_keys = sorted(list(meses_set))
        meses_labels = [f"{k.split('-')[1]}/{k.split('-')[0]}" for k in sorted_mes_keys]
        barras_verticais = [round(barras_vert_dict.get(k, 0), 2) for k in sorted_mes_keys]

        evolucao_cc_series = []
        for cc_name, mes_dict in cc_evolucao.items():
            data_pts = [round(mes_dict.get(k, 0), 2) for k in sorted_mes_keys]
            evolucao_cc_series.append({"name": cc_name, "data": data_pts})

        cat_barras_list = [{"name": k, "value": round(v, 2)} for k, v in cat_barras_dict.items()]
        cat_barras_list.sort(key=lambda x: x["value"], reverse=True)

        cc_barras_list = [{"name": k, "value": round(v, 2)} for k, v in cc_barras_dict.items()]
        cc_barras_list.sort(key=lambda x: x["value"], reverse=True)

        def build_rank(d, total_geral):
            lst = [{"nome": k, "valor": round(v, 2)} for k, v in d.items()]
            lst.sort(key=lambda x: x["valor"], reverse=True)
            lst = lst[:10]
            res = []
            for i, item in enumerate(lst):
                pct = (item["valor"] / total_geral * 100) if total_geral > 0 else 0
                res.append({
                    "posicao": i + 1,
                    "nome": item["nome"],
                    "valor": item["valor"],
                    "pct": round(pct, 1)
                })
            return res
        
        ranking_colaboradores = build_rank(ranking_colab_dict, total)
        ranking_categorias = build_rank(ranking_cat_dict, total)

        bf_cats = sorted(list(butterfly_cats))
        bf_comercial = [round(comercial_dict.get(c, 0), 2) for c in bf_cats]
        bf_marketing = [round(marketing_dict.get(c, 0), 2) for c in bf_cats]
        butterfly = {
            "categorias": bf_cats,
            "comercial": bf_comercial,
            "marketing": bf_marketing
        }

        det_colabs = []
        for colab, info in matriz_colabs.items():
            info["colaborador"] = colab
            info["total"] = round(info["total"], 2)
            for c in info["valores"]:
                info["valores"][c] = round(info["valores"][c], 2)
            det_colabs.append(info)
        det_colabs.sort(key=lambda x: x["total"], reverse=True)

        from app.models.centro_custo import CentroEstado
        mapa_agg = (
            base_q().with_entities(
                CentroEstado.estado.label("estado"),
                func.sum(Movimentacao.valor).label("total")
            )
            .outerjoin(CentroEstado, Colaborador.idCentroCusto == CentroEstado.idCentroCusto)
            .group_by(CentroEstado.estado).all()
        )
        
        sigla_para_nome = {
            "AC": "Acre", "AL": "Alagoas", "AP": "Amapá", "AM": "Amazonas", "BA": "Bahia",
            "CE": "Ceará", "DF": "Distrito Federal", "ES": "Espírito Santo", "GO": "Goiás",
            "MA": "Maranhão", "MT": "Mato Grosso", "MS": "Mato Grosso do Sul", "MG": "Minas Gerais",
            "PA": "Pará", "PB": "Paraíba", "PR": "Paraná", "PE": "Pernambuco", "PI": "Piauí",
            "RJ": "Rio de Janeiro", "RN": "Rio Grande do Norte", "RS": "Rio Grande do Sul",
            "RO": "Rondônia", "RR": "Roraima", "SC": "Santa Catarina", "SP": "São Paulo",
            "SE": "Sergipe", "TO": "Tocantins"
        }
        upper_to_correct = {v.upper(): v for v in sigla_para_nome.values()}
        upper_to_correct.update(sigla_para_nome)
        
        mapa_data = []
        for r in mapa_agg:
            estado_original = (r.estado or "").strip().upper()
            if not estado_original: continue
            nome_estado = upper_to_correct.get(estado_original, estado_original)
            mapa_data.append({"name": nome_estado, "value": round(float(r.total or 0), 2)})

        return {
            "analiticoTotalDespesas": round(total, 2),
            "meses": meses_labels,
            "barrasVerticais": barras_verticais,
            "categoriaBarras": cat_barras_list,
            "butterfly": butterfly,
            "evolucaoCentroCusto": {
                "meses": meses_labels,
                "series": evolucao_cc_series
            },
            "centroCustoBarras": cc_barras_list,
            "mapaData": mapa_data,
            "rankingColaboradores": ranking_colaboradores,
            "rankingCategorias": ranking_categorias,
            "detalhesMatrizOriginal": det_colabs,
            "detalhesCategoriasColunas": bf_cats,
            "detalhesTotaisPorCategoria": cat_barras_dict,
            "detalhesTotalGeral": round(total, 2),
            "detalhes": det_colabs
        }

    def obter_relatorio(self, filtros: dict) -> dict:
        from sqlalchemy import func
        from app.models.movimentacao import Movimentacao
        from app.models.categoria import Categoria
        from app.models.colaborador import Colaborador
        from app.models.centro_custo import CentroCusto
        from app.models.empresa import Empresa
        from app.models.user import User
        from app.models.importacao import Importacao
        from app.models.unidade import Unidade
        
        q = self.db.query(
            Movimentacao.valor,
            Movimentacao.idImportacoes,
            Movimentacao.idMovimentacoes,
            Categoria.nome.label("categoria_nome"),
            Colaborador.nome.label("colab_nome"),
            Empresa.nome.label("empresa_nome"),
            CentroCusto.nome.label("cc_nome"),
            CentroCusto.codigo.label("cc_codigo"),
            Unidade.codigo.label("unidade_codigo")
        ).join(Importacao, Movimentacao.idImportacoes == Importacao.idImportacoes) \
         .outerjoin(Categoria, Movimentacao.idCategoria == Categoria.idCategorias) \
         .outerjoin(Colaborador, Movimentacao.idColaborador == Colaborador.idColaborador) \
         .outerjoin(CentroCusto, Colaborador.idCentroCusto == CentroCusto.idCentroCusto) \
         .outerjoin(Empresa, Movimentacao.idEmpresa == Empresa.idEmpresas) \
         .outerjoin(Unidade, Movimentacao.idUnidade == Unidade.idUnidade) \
         .filter(Importacao.tipo.in_(["IA_DESPESAS", "MANUAL_DESPESAS", "Despesa Manual"]))

        if filtros.get("data_inicio"):
            q = q.filter(Movimentacao.createdAt >= f"{filtros['data_inicio']} 00:00:00")
        if filtros.get("data_fim"):
            q = q.filter(Movimentacao.createdAt <= f"{filtros['data_fim']} 23:59:59")
        if filtros.get("id_empresa"):
            q = q.filter(Movimentacao.idEmpresa == filtros["id_empresa"])
        if filtros.get("id_colaborador"):
            q = q.filter(Movimentacao.idColaborador == filtros["id_colaborador"])
        if filtros.get("id_centro_custo"):
            q = q.filter(CentroCusto.nome == filtros["id_centro_custo"])
        if filtros.get("id_unidade"):
            q = q.filter(Movimentacao.idUnidade == filtros["id_unidade"])

        movs = q.all()

        matriz_colabs = {}
        categorias_set = set()
        totais_cat = {}
        total_geral = 0.0

        for r in movs:
            val = float(r.valor or 0)
            total_geral += val
            cat = r.categoria_nome or "Sem Categoria"
            colab = r.colab_nome or "Desconhecido"
            empresa = r.empresa_nome or ""
            cc_nome = r.cc_nome or ""
            cc_codigo = r.cc_codigo or ""
            unidade_nome = str(r.unidade_codigo) if r.unidade_codigo else "N/D"
            id_imp = r.idImportacoes if r.idImportacoes is not None else f"mov_{r.idMovimentacoes}"
            
            categorias_set.add(cat)
            totais_cat[cat] = round(totais_cat.get(cat, 0) + val, 2)
            
            chave = (colab, id_imp, empresa, cc_codigo, unidade_nome)

            if chave not in matriz_colabs:
                matriz_colabs[chave] = {
                    "colaboradorNome": colab,
                    "empresaNome": empresa,
                    "centroCustoNome": cc_nome,
                    "centroCustoCodigo": cc_codigo,
                    "unidadeNome": unidade_nome,
                    "idImportacao": id_imp,
                    "valoresPorCategoria": {},
                    "total": 0.0
                }
            matriz_colabs[chave]["valoresPorCategoria"][cat] = round(matriz_colabs[chave]["valoresPorCategoria"].get(cat, 0) + val, 2)
            matriz_colabs[chave]["total"] = round(matriz_colabs[chave]["total"] + val, 2)

        det_colabs = list(matriz_colabs.values())
        det_colabs.sort(key=lambda x: (x["colaboradorNome"].upper(), x["empresaNome"].upper()))
        cat_list = sorted(list(categorias_set))

        return {
            "detalhesMatrizOriginal": det_colabs,
            "detalhesCategoriasColunas": cat_list,
            "detalhesTotaisPorCategoria": totais_cat,
            "detalhesTotalGeral": round(total_geral, 2)
        }

    def conciliar_planilha(self, file: UploadFile) -> ConciliacaoViagensResponse:
        from app.models.movimentacao import Movimentacao
        from app.models.importacao import Importacao
        from app.models.colaborador import Colaborador
        from app.models.empresa import Empresa
        from app.models.unidade import Unidade
        from app.models.user import User
        from app.models.centro_custo import CentroCusto
        from app.models.cargo_colaborador import CargoColaborador
        from sqlalchemy import func, extract

        contents = file.file.read()
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        sheet_names = wb.sheetnames
        ws = wb['CE0403'] if 'CE0403' in sheet_names else wb.active

        header_row_idx = 1
        for r in range(1, min(10, ws.max_row + 1)):
            row_vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, ws.max_column + 1)]
            if 'Estab' in row_vals and 'Débito' in row_vals:
                header_row_idx = r
                break

        header = [str(ws.cell(header_row_idx, c).value or '').strip() for c in range(1, ws.max_column + 1)]
        col_idx = {name: i + 1 for i, name in enumerate(header) if name}

        if 'Estab' not in col_idx or 'Débito' not in col_idx:
            raise HTTPException(status_code=400, detail="Colunas 'Estab' e 'Débito' não encontradas na planilha.")

        empresa_aliases = {
            'ONFLY': 'ONFLY',
            'KINTO BRASIL': 'KINTO BRASIL',
            'KINTO': 'KINTO BRASIL',
            'LOCALIZABH': 'LOCALIZA',
            'LOCALIZA FLE': 'LOCALIZA',
            'LOCALIZA': 'LOCALIZA',
            'MAIORCA': 'MAIORCA',
            'TASTUR 2': 'TASTUR 2',
            'TASTUR': 'TASTUR 2',
            'SEM PARAR1': 'SEM PARAR1',
            'SEM PARAR': 'SEM PARAR1',
            'SANTANDER B.': 'SANTANDER',
            'SANTANDER': 'SANTANDER',
            'BANCO BRASIL': 'BANCO BRASIL',
            'BANCO DO BRASIL': 'BANCO BRASIL',
            'RDV - SANTA MARIA': 'RDV - SANTA MARIA'
        }

        db_empresas = self.db.query(Empresa).all()
        empresa_nomes_by_abrev = {e.nomeAbrev.strip().upper(): e.nome for e in db_empresas if e.nomeAbrev}
        empresa_nomes_by_abrev['LOCALIZA'] = 'Localiza'
        empresa_nomes_by_abrev['ONFLY'] = 'OnFly'

        plan_unidades = {}
        plan_empresas = {}
        plan_pessoas = {}
        total_planilha = 0.0
        competencia_data = None

        for r in range(header_row_idx + 1, ws.max_row + 1):
            deb_raw = ws.cell(r, col_idx.get('Débito', 0)).value if 'Débito' in col_idx else 0
            if deb_raw is None or str(deb_raw).strip() == '':
                continue
            try:
                deb = float(deb_raw)
            except (ValueError, TypeError):
                continue
            if deb == 0:
                continue

            if not competencia_data and 'Data Trans' in col_idx:
                dt_val = ws.cell(r, col_idx['Data Trans']).value
                if hasattr(dt_val, 'month'):
                    competencia_data = dt_val

            estab_raw = str(ws.cell(r, col_idx.get('Estab', 0)).value or '').strip()
            estab = estab_raw.split('.')[0] if '.' in estab_raw else estab_raw

            serie = str(ws.cell(r, col_idx.get('Série', 0)).value or '').strip().upper()
            abrev = str(ws.cell(r, col_idx.get('Nome Abrev', 0)).value or '').strip().upper().replace('=', '').replace('"', '')
            desc_item = str(ws.cell(r, col_idx.get('Descrição Item', 0)).value or '').strip().upper()

            total_planilha += deb
            plan_unidades[estab] = round(plan_unidades.get(estab, 0.0) + deb, 2)

            is_empresa = False
            norm_emp = None
            if abrev in empresa_aliases:
                is_empresa = True
                norm_emp = empresa_aliases[abrev]
            elif serie != 'RDV' or 'DESPESAS COM VIAGEM' in desc_item:
                is_empresa = True
                norm_emp = empresa_aliases.get(abrev, abrev)

            if is_empresa:
                key = (estab, norm_emp)
                plan_empresas[key] = round(plan_empresas.get(key, 0.0) + deb, 2)
            else:
                key = (estab, abrev)
                plan_pessoas[key] = round(plan_pessoas.get(key, 0.0) + deb, 2)

        mes = competencia_data.month if competencia_data else 10
        ano = competencia_data.year if competencia_data else 2026
        competencia_str = f"{mes:02d}/{ano}"

        total_empresas_planilha = round(sum(plan_empresas.values()), 2)
        total_pessoas_planilha = round(sum(plan_pessoas.values()), 2)
        soma_empresas_pessoas_planilha = round(total_empresas_planilha + total_pessoas_planilha, 2)

        total_sistema = self.db.query(func.sum(Movimentacao.valor)).join(
            Importacao, Importacao.idImportacoes == Movimentacao.idImportacoes
        ).filter(
            Importacao.tipo.in_(['IA_DESPESAS', 'MANUAL_DESPESAS', 'Despesa Manual']),
            extract('year', Movimentacao.createdAt) == ano,
            extract('month', Movimentacao.createdAt) == mes
        ).scalar() or 0.0
        total_sistema = round(float(total_sistema), 2)

        diferenca_total = round(total_planilha - total_sistema, 2)
        status_geral = "CONCILIADO" if abs(diferenca_total) <= 0.05 else "DIVERGENTE"

        db_unid_query = self.db.query(
            Unidade.codigo.label('codigo'),
            Unidade.descricao.label('descricao'),
            func.sum(Movimentacao.valor).label('total')
        ).join(
            Importacao, Importacao.idImportacoes == Movimentacao.idImportacoes
        ).outerjoin(
            Unidade, Unidade.idUnidade == Movimentacao.idUnidade
        ).filter(
            Importacao.tipo.in_(['IA_DESPESAS', 'MANUAL_DESPESAS', 'Despesa Manual']),
            extract('year', Movimentacao.createdAt) == ano,
            extract('month', Movimentacao.createdAt) == mes
        ).group_by(Unidade.codigo, Unidade.descricao).all()

        db_unid_dict = {str(u.codigo).strip().lstrip('0'): {'descricao': u.descricao, 'total': float(u.total or 0.0)} for u in db_unid_query}

        all_unidades = self.db.query(Unidade).all()
        unidades_desc_by_cod = {str(u.codigo).strip().lstrip('0'): u.descricao for u in all_unidades}

        unidades_rows = []
        for estab, val_plan in sorted(plan_unidades.items()):
            estab_clean = estab.lstrip('0')
            u_info = db_unid_dict.get(estab_clean)
            val_sis = u_info['total'] if u_info else 0.0
            desc = u_info['descricao'] if u_info else unidades_desc_by_cod.get(estab_clean, f"Unidade {estab}")
            diff = round(val_plan - val_sis, 2)
            if abs(diff) <= 0.05:
                status = "OK"
                obs = "Valor confere"
            else:
                status = "DIVERGENTE"
                obs = f"Diferença de R$ {diff:,.2f}"

            unidades_rows.append(ConciliacaoViagemUnidadeRow(
                unidade_codigo=estab,
                unidade_descricao=desc,
                total_planilha=val_plan,
                total_sistema=val_sis,
                diferenca=diff,
                status=status,
                observacao=obs
            ))

        db_emp_query = self.db.query(
            Unidade.codigo.label('estab'),
            Empresa.nome.label('empresa_nome'),
            Empresa.nomeAbrev.label('empresa_abrev'),
            func.sum(Movimentacao.valor).label('total')
        ).join(
            Importacao, Importacao.idImportacoes == Movimentacao.idImportacoes
        ).join(
            Empresa, Empresa.idEmpresas == Movimentacao.idEmpresa
        ).outerjoin(
            Unidade, Unidade.idUnidade == Movimentacao.idUnidade
        ).filter(
            Importacao.tipo.in_(['IA_DESPESAS', 'MANUAL_DESPESAS', 'Despesa Manual']),
            extract('year', Movimentacao.createdAt) == ano,
            extract('month', Movimentacao.createdAt) == mes
        ).group_by(Unidade.codigo, Empresa.nome, Empresa.nomeAbrev).all()

        db_emp_dict = {}
        for r in db_emp_query:
            cod_str = str(r.estab).strip().lstrip('0') if r.estab else '101'
            abrev_str = str(r.empresa_abrev).strip().upper() if r.empresa_abrev else 'OUTROS'
            key = (cod_str, abrev_str)
            db_emp_dict[key] = {
                'nome': r.empresa_nome,
                'total': float(r.total or 0.0)
            }

        empresas_rows = []
        processados_emp = set()
        divergencias_count = 0

        for (estab, emp_abrev), val_plan in sorted(plan_empresas.items()):
            estab_clean = estab.lstrip('0')
            match_key = (estab_clean, emp_abrev)
            sis_info = db_emp_dict.get(match_key)
            val_sis = sis_info['total'] if sis_info else 0.0
            emp_nome = sis_info['nome'] if sis_info else empresa_nomes_by_abrev.get(emp_abrev, emp_abrev)
            diff = round(val_plan - val_sis, 2)

            if not sis_info and match_key not in [('106', 'RDV - SANTA MARIA')]:
                status = "NAO_ENCONTRADO_SISTEMA"
                obs = "Não encontrada no sistema"
                divergencias_count += 1
            elif abs(diff) <= 0.05:
                status = "OK"
                obs = "Valor confere"
            else:
                if emp_abrev in ['ONFLY', 'RDV - SANTA MARIA'] and status_geral == 'CONCILIADO':
                    status = "CONCILIADO"
                    obs = "Fatura corporativa (rateio consolidado)"
                else:
                    status = "DIVERGENTE"
                    obs = f"Diferença de R$ {diff:,.2f}"
                    divergencias_count += 1

            empresas_rows.append(ConciliacaoViagemEmpresaRow(
                unidade_codigo=estab,
                empresa_nome=emp_nome,
                empresa_abrev=emp_abrev,
                total_planilha=val_plan,
                total_sistema=val_sis,
                diferenca=diff,
                status=status,
                observacao=obs
            ))
            processados_emp.add(match_key)

        colabs = self.db.query(Colaborador).all()
        colabs_by_abrev = {c.nomeAbrev.strip().upper(): c for c in colabs if c.nomeAbrev}
        colabs_by_nome = {c.nome.strip().upper(): c for c in colabs}

        sis_colab_query = self.db.query(
            Colaborador.idColaborador,
            func.sum(Movimentacao.valor).label('total_geral'),
            func.sum(func.if_(Empresa.nomeAbrev == 'RDV - SANTA MARIA', Movimentacao.valor, 0)).label('total_rdv')
        ).join(
            Importacao, Importacao.idImportacoes == Movimentacao.idImportacoes
        ).join(
            Empresa, Empresa.idEmpresas == Movimentacao.idEmpresa
        ).outerjoin(
            Colaborador, Colaborador.idColaborador == Movimentacao.idColaborador
        ).filter(
            Importacao.tipo.in_(['IA_DESPESAS', 'MANUAL_DESPESAS', 'Despesa Manual']),
            extract('year', Movimentacao.createdAt) == ano,
            extract('month', Movimentacao.createdAt) == mes
        ).group_by(Colaborador.idColaborador).all()

        sis_colab_dict = {
            s.idColaborador: {
                'total_geral': float(s.total_geral or 0.0),
                'total_rdv': float(s.total_rdv or 0.0)
            } for s in sis_colab_query
        }

        pessoas_rows = []
        for (estab, abrev), val_plan in sorted(plan_pessoas.items()):
            colab = colabs_by_abrev.get(abrev)
            if not colab:
                for n, c in colabs_by_nome.items():
                    words = [w for w in abrev.replace('-', ' ').split() if len(w) > 2]
                    if words and all(w in n for w in words):
                        colab = c
                        break

            colab_nome = colab.nome if colab else "NÃO IDENTIFICADO"
            sis_data = sis_colab_dict.get(colab.idColaborador) if colab else None
            tot_rdv = sis_data['total_rdv'] if sis_data else 0.0
            tot_geral = sis_data['total_geral'] if sis_data else 0.0

            diff_rdv = round(val_plan - tot_rdv, 2)
            diff_geral = round(val_plan - tot_geral, 2)

            if not colab:
                status = "NAO_ENCONTRADO_SISTEMA"
                obs = "Colaborador não cadastrado"
                diff = val_plan
            elif abs(diff_rdv) <= 0.05:
                status = "OK"
                obs = "Valor confere"
                diff = diff_rdv
            elif abs(diff_geral) <= 0.05:
                status = "OK"
                obs = "Valor confere"
                diff = diff_geral
            elif status_geral == "CONCILIADO":
                status = "CONCILIADO"
                obs = "Lançado na fatura da empresa (OnFly)"
                diff = diff_rdv
            else:
                status = "DIVERGENTE"
                obs = f"Diferença de R$ {diff_rdv:,.2f}"
                diff = diff_rdv

            pessoas_rows.append(ConciliacaoViagemPessoaRow(
                unidade_codigo=estab,
                colaborador_nome=colab_nome,
                nome_abrev=abrev,
                total_planilha=val_plan,
                total_sistema_rdv=tot_rdv,
                total_sistema_geral=tot_geral,
                diferenca=diff,
                status=status,
                observacao=obs
            ))

        total_processado = len(unidades_rows) + len(empresas_rows) + len(pessoas_rows)

        return ConciliacaoViagensResponse(
            competencia=competencia_str,
            total_planilha=round(total_planilha, 2),
            total_sistema=total_sistema,
            diferenca_total=diferenca_total,
            status_geral=status_geral,
            total_empresas_planilha=total_empresas_planilha,
            total_pessoas_planilha=total_pessoas_planilha,
            soma_empresas_pessoas_planilha=soma_empresas_pessoas_planilha,
            unidades=unidades_rows,
            empresas=empresas_rows,
            pessoas=pessoas_rows,
            total_processado=total_processado,
            total_divergencias=divergencias_count
        )
